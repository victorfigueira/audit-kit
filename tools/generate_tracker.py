#!/usr/bin/env python3
"""Create or update audit/audit-tracker.xlsx from audit/findings.json.

Usage: python3 audit-kit/tools/generate_tracker.py [findings.json] [tracker.xlsx]

Manual columns (Status, Owner, Target date, Closed date, Jira key, Notes) are
preserved across runs. Findings missing from the latest audit are kept and
marked "Not found in latest audit".
"""
import json
import os
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is missing. Run: pip3 install openpyxl")

DIMENSIONS = [
    "Maintainability", "Open Source Health", "Architecture", "Security",
    "Reliability", "Performance and Scalability", "Testing and QA",
    "Observability and Operability", "Data and Privacy Compliance",
    "Infra, CI/CD and Cloud Cost",
]
SEVERITIES = ["Critical", "High", "Medium", "Low"]
STATUSES = ["Open", "In progress", "Fixed", "Accepted risk",
            "Not found in latest audit"]
COLUMNS = ["ID", "Dimension", "Severity", "Title", "Location",
           "Recommendation", "Effort", "Status", "Owner", "Target date",
           "Closed date", "Jira key", "Notes"]
MANUAL = ["Status", "Owner", "Target date", "Closed date", "Jira key", "Notes"]
WIDTHS = [11, 28, 11, 45, 35, 55, 8, 22, 16, 13, 13, 12, 40]
NOT_FOUND = "Not found in latest audit"

FONT = "Arial"
NAVY = PatternFill("solid", fgColor="1F3864")
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
BODY = Font(name=FONT)
WRAP = Alignment(wrap_text=True, vertical="top")


def read_existing(path):
    """Return {key: row dict} from an existing Findings sheet."""
    rows = {}
    if not os.path.exists(path):
        return rows
    wb = load_workbook(path)
    if "Findings" not in wb.sheetnames:
        return rows
    ws = wb["Findings"]
    header = [c.value for c in ws[1]]
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, values))
        if row.get("ID"):
            rows[row["ID"]] = row
    return rows


def style_header(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font, cell.alignment = NAVY, HEAD, WRAP
    ws.freeze_panes = "A2"


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "audit/findings.json"
    dst = sys.argv[2] if len(sys.argv) > 2 else "audit/audit-tracker.xlsx"
    with open(src, encoding="utf-8") as fh:
        data = json.load(fh)

    old = read_existing(dst)
    old_by_title = {(r.get("Title"), r.get("Location")): r for r in old.values()}

    rows, seen = [], set()
    for f in data.get("findings", []):
        prev = old.get(f["id"]) or old_by_title.get((f.get("title"), f.get("location")))
        row = {
            "ID": f["id"], "Dimension": f.get("dimension"),
            "Severity": f.get("severity"), "Title": f.get("title"),
            "Location": f.get("location"), "Recommendation": f.get("recommendation"),
            "Effort": f.get("effort"), "Status": "Open",
        }
        if prev:
            seen.add(prev["ID"])
            for k in MANUAL:
                if prev.get(k) not in (None, ""):
                    row[k] = prev[k]
            if row["Status"] == NOT_FOUND:
                row["Status"] = "Open"
        rows.append(row)
    for key, prev in old.items():
        if key not in seen and key not in {r["ID"] for r in rows}:
            prev = dict(prev)
            if prev.get("Status") not in ("Fixed", "Accepted risk"):
                prev["Status"] = NOT_FOUND
            rows.append(prev)

    sev_order = {s: i for i, s in enumerate(SEVERITIES)}
    rows.sort(key=lambda r: (r.get("Status") == NOT_FOUND,
                             sev_order.get(r.get("Severity"), 9), r["ID"]))

    wb = Workbook()
    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Audit tracker – {data.get('project', '')}"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = (f"Mode: {data.get('mode', '')} · Branch: {data.get('branch', '')} · "
                f"Commit: {data.get('commit', '')} · Date: {data.get('date', '')}")
    ws["A2"].font = BODY
    head = ["Dimension", "Score"] + SEVERITIES + ["Total", "Open", "Fixed"]
    for i, h in enumerate(head, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.fill, c.font, c.alignment = NAVY, HEAD, WRAP
    n = len(rows) + 1
    rng = lambda col: f"Findings!${col}$2:${col}${max(n, 2)}"
    scores = data.get("scores", {})
    first = 5
    for r, dim in enumerate(DIMENSIONS, first):
        ws.cell(row=r, column=1, value=dim)
        ws.cell(row=r, column=2, value=scores.get(dim, "Not audited"))
        for j, sev in enumerate(SEVERITIES, 3):
            ws.cell(row=r, column=j,
                    value=f'=COUNTIFS({rng("B")},$A{r},{rng("C")},"{sev}",{rng("H")},"<>{NOT_FOUND}")')
        ws.cell(row=r, column=7, value=f"=SUM(C{r}:F{r})")
        ws.cell(row=r, column=8, value=f'=COUNTIFS({rng("B")},$A{r},{rng("H")},"Open")')
        ws.cell(row=r, column=9, value=f'=COUNTIFS({rng("B")},$A{r},{rng("H")},"Fixed")')
    last = first + len(DIMENSIONS) - 1
    tot = last + 1
    ws.cell(row=tot, column=1, value="Total / overall score").font = Font(name=FONT, bold=True)
    ws.cell(row=tot, column=2, value=f'=IFERROR(ROUND(AVERAGE(B{first}:B{last}),1),"-")')
    for col in "CDEFGHI":
        ws[f"{col}{tot}"] = f"=SUM({col}{first}:{col}{last})"
    for row in ws.iter_rows(min_row=first, max_row=tot):
        for c in row:
            c.font = Font(name=FONT, bold=(c.row == tot))
    ws.column_dimensions["A"].width = 32
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 12

    # Findings
    wf = wb.create_sheet("Findings")
    wf.append(COLUMNS)
    style_header(wf, len(COLUMNS))
    for row in rows:
        wf.append([row.get(c) for c in COLUMNS])
    for i, w in enumerate(WIDTHS, 1):
        wf.column_dimensions[wf.cell(row=1, column=i).column_letter].width = w
    for r in wf.iter_rows(min_row=2):
        for c in r:
            c.font, c.alignment = BODY, WRAP
    wf.auto_filter.ref = f"A1:M{max(n, 2)}"
    end = max(n, 2) + 500
    for col, src_col, values in (("B", "A", DIMENSIONS), ("C", "B", SEVERITIES),
                                 ("H", "C", STATUSES)):
        dv = DataValidation(type="list", allow_blank=True,
                            formula1=f"=Legend!${src_col}$2:${src_col}${len(values) + 1}")
        wf.add_data_validation(dv)
        dv.add(f"{col}2:{col}{end}")

    # Legend
    wl = wb.create_sheet("Legend")
    wl.append(["Dimension", "Severity", "Status", "", "Severity meaning", "", "Score", "Meaning"])
    style_header(wl, 8)
    meaning = ["Exploitable, data loss, legal exposure or outage likely. Fix immediately.",
               "Serious risk or blocker. Fix this sprint.",
               "Limited impact or workaround. Plan this quarter.",
               "Improvement or hygiene. Backlog."]
    scale = [("9–10", "Exemplary, no High or Critical"), ("7–8", "Solid, a few Medium"),
             ("5–6", "Acceptable, at least one High"), ("3–4", "Weak, several High or one Critical"),
             ("1–2", "Severe, multiple Critical")]
    for i in range(max(len(DIMENSIONS), len(STATUSES))):
        r = i + 2
        if i < len(DIMENSIONS): wl.cell(row=r, column=1, value=DIMENSIONS[i])
        if i < len(SEVERITIES):
            wl.cell(row=r, column=2, value=SEVERITIES[i])
            wl.cell(row=r, column=5, value=f"{SEVERITIES[i]}: {meaning[i]}")
        if i < len(STATUSES): wl.cell(row=r, column=3, value=STATUSES[i])
        if i < len(scale):
            wl.cell(row=r, column=7, value=scale[i][0])
            wl.cell(row=r, column=8, value=scale[i][1])
    wl.cell(row=14, column=1, value="Edit only Status, Owner, Target date, Closed date, Jira key and Notes in Findings. Other columns are rewritten on every audit.")
    for row in wl.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
    for col, w in zip("ABCDEFGH", [32, 12, 26, 3, 70, 3, 8, 40]):
        wl.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(dst) or ".", exist_ok=True)
    wb.save(dst)
    print(f"Tracker written: {dst} ({len(rows)} findings)")


if __name__ == "__main__":
    main()
